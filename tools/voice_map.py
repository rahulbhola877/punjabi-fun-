"""
Voice map: which clip may be said by whom, when, and from how far away.

  python3 tools/voice_map.py make     -> writes docs/voice-map.xlsx (fill it in) and src/audio/voiceMap.json (defaults)
  python3 tools/voice_map.py import F -> reads a filled-in workbook F and rewrites src/audio/voiceMap.json

The game reads only src/audio/voiceMap.json.
"""
import json, sys
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

VOICES = {
    "A": "jolly man (Bhola's voice)", "B": "dry woman (Gurnam's voice)", "C": "cheerful woman (Pammi's voice)",
    "D": "gruff man (police voice)", "E": "young man (lafanga voice)", "F": "loud man (angry neighbour voice)",
    "G": "old man (Ustad's voice)", "?": "made by the other session", "S": "Punjabi shopkeeper man (Saurav)", "P": "the player, male (Deepak)", "Q": "the player, female (Paaro)", "R": "the player's shout, desi Punjabi (Pind Waali)",
    "T": "aunty, middle-aged woman (Ritu Rathore)", "U": "young woman, Hinglish (Pyaari)",
}

# Who codes: named characters (by id) and groups.
WHO = [
    ("bhola", "Bhola Singh — Bhole Da Dhaba, Ranggarh bazaar"), ("cook", "Kaku the cook — dhaba kitchen"),
    ("gurnam", "Gurnam Kaur — mechanic, Gurnam Motor Works"), ("pammi", "Pammi — fruit & sabzi stall"),
    ("masterji", "Masterji — tailor"), ("sohal", "Sohal — Sohal Cycle Store"), ("mrsKaur", "Mrs. Kaur — elderly neighbour, courtyard"),
    ("gill", "Gill uncle — lane house"), ("chemist", "Dr. Sehgal — Sehat Chemist"), ("mobileWala", "Pinka — Hello Mobile Zone, Chaura Bazaar"),
    ("ustad", "Ustad Kartar Pehlwan — Kotli akhada"), ("conductor", "Conductor Bittu — bus adda"), ("munshi", "Munshi Ji — anaj mandi"),
    ("balwant", "Havaldar Balwant — police chowki"), ("ticketwala", "Ticket-wala Raju — Raj Talkies"), ("sukhi", "Sukhi Bhaji — trucker at Sher-e-Punjab dhaba"),
    ("sherOwner", "Sher Singh — Sher-e-Punjab dhaba"), ("sarpanch", "Sarpanch Hardev Singh — Pind Kotli"), ("jarnail", "Jarnail Singh — dera"),
    ("pumpwala", "Pump attendant — petrol pump"), ("clerk_bank", "Anand sahib — bank manager"), ("clerk_post", "Dhillon ji — postmaster"),
    ("any-man", "any adult man"), ("young-man", "a young man"), ("any-woman", "any adult woman"), ("any-elder", "any old man or woman"),
    ("shopkeeper-man", "a male shopkeeper"), ("shopkeeper-woman", "a female shopkeeper"), ("aunty", "an aunty (in an aunty quarrel)"),
    ("police", "any constable / havaldar"), ("police-jeep", "the police jeep's loudspeaker"), ("police-bike", "a police motorbike rider"), ("player", "you (the player character)"), ("lafanga", "a lafanga (street tough)"),
    ("auto-wala", "an auto-wala at a stand"), ("bus-conductor", "a bus conductor"), ("sabzi-wala", "sabzi seller"), ("fruit-wala", "fruit seller"),
    ("chaat-wala", "gol-gappe / chaat seller"), ("doodh-wala", "milkman"), ("kabadi-wala", "scrap buyer"),
    ("driver", "a driver you bumped"), ("passenger", "an auto passenger"), ("onlooker", "a bystander watching a fight"), ("anyone", "anyone at all"),
]
WHEN = [
    ("talk", "you talk to them (dialogue box opens)"), ("mission-start", "first mission: Bhola sends you to Chaura Bazaar"),
    ("mission-shop", "first mission: buying the phone"), ("mission-end", "first mission: giving Bhola the phone"),
    ("pass-by", "you walk past them (within the distance)"), ("chat", "you pass two people chatting"), ("street-call", "a seller/tout calls out when you are near"),
    ("bumped", "you walk or drive into them"), ("hit", "you punch or slap them"), ("police-chase", "constables start chasing you"),
    ("police-bribe", "police catch you (chai-paani)"), ("police-jeep", "the police jeep chases you"), ("lafanga-taunt", "a lafanga taunts you"),
    ("gang-fight", "a gang picks a fight with you"), ("quarrel-start", "two men start a street quarrel"), ("quarrel-blows", "a men's quarrel comes to blows"),
    ("aunty-quarrel-start", "two aunties start a quarrel"), ("aunty-quarrel-blows", "an aunty quarrel comes to blows"),
    ("quarrel-onlooker", "a bystander comes to watch a quarrel"), ("blame-both", "you tell both sides they are wrong"),
    ("road-rage", "a driver gets out after you bump his car"), ("safari-hit", "you hit the Safari"), ("auto-tipped", "an overloaded auto tips over"),
    ("ride-jolt", "your auto passenger gets a bad jolt"), ("dangal-start", "the akhada dangal starts"),
    ("fight-start", "someone squares up to fight you"), ("fight-onlooker", "a bystander reacts to your fight"),
    ("meet", "you stop next to someone (your greeting)"), ("meet-reply", "they answer your greeting"),
    ("colony-fight-start", "a group fight starts in a residential colony"), ("colony-fight-blows", "the colony fight comes to blows"),
    ("crash-drivers", "two vehicles crash and the drivers get out"), ("police-patrol", "a police jeep/bike passes on patrol"), ("snake-hit", "you knock a zig-zag snake biker off his bike"),
    ("fight-elder", "you square up to an old uncle"), ("rain-start", "it starts to pour: someone near you reacts"), ("flood", "the galis are flooded: someone near you complains"),
    ("doorbell-stranger", "you ring a stranger's bell and they answer at the gate"), ("doorbell-family", "you ring the bell of a family that knows yours"),
    ("doorbell-night", "you ring a bell at night: a voice from a window"), ("doorbell-siesta", "you ring a bell in the afternoon rest hour"),
    ("power-cut", "the light goes: someone near you complains"), ("power-back", "the light comes back: someone near you cheers"),
    ("rail-gate", "the level-crossing gates close: someone near you complains"), ("pump-arrive", "you pull up at the petrol pump"), ("pump-done", "the tank is full"),
    ("fuel-out", "your tank runs dry"), ("bank-welcome", "you step up to the bank counter"), ("bank-advice", "you finish at the bank counter"), ("post-welcome", "you collect the post"),
    ("signal-jump", "you cross a junction against a red light: a bystander shouts"), ("meet-friend", "a friend answers your greeting"), ("vehicle-admire", "you get off a vehicle and someone near you admires it"),
]

# clip, words, voice, who (codes), when (codes), max distance m, min gap s, on
CLIPS = [
    ("bhola_hello", "Oye! Sat Sri Akal ji! Hot food waits for no one — go, go, go!", "A", "bhola", "talk", 12, 60, "Yes"),
    ("bhola_chaura", "Oye! Chaure Bazaar jaana ae, mobile lain layi! Mera phone thalle digg ke tutt gaya — screen de sau tukde!", "A", "bhola", "mission-start", 15, 0, "Yes"),
    ("bhola_thanks", "Wah mere sher! Hun main vi status paavanga! Eh lai, tera inaam.", "A", "bhola", "mission-end", 15, 0, "Yes"),
    ("mobile_shop", "Aao ji, aao! Naya model, kamaal da camera! Bhole da phone? Eh lo, pack kar ditta!", "E", "mobileWala", "mission-shop", 15, 0, "Yes"),
    ("gurnam_finally", "Finally. Bhola's dal makhani. The only fuel I trust.", "B", "gurnam", "talk", 12, 120, "No"),
    ("pammi_kinnow", "Oho! Kinnow from Abohar, sweet as gur! And the gossip comes free!", "C", "pammi", "pass-by, talk", 8, 120, "Yes"),
    ("masterji_naap", "Naap do vaari lao, kainchi ik vaari! Tera kurta vadda ae, puttar.", "G", "masterji", "pass-by", 6, 120, "Yes"),
    ("sohal_race", "Ik din tainu gate tak race vich harauna ae, dekh layi!", "E", "sohal", "pass-by", 6, 120, "Yes"),
    ("mrskaur_chah", "Aa ja puttar, chah pee ke ja! Bina khaade kithe challeya?", "B", "mrsKaur", "pass-by", 6, 120, "Yes"),
    ("gill_vele", "Saade vele gali kachi si, par dil pakke si!", "G", "gill", "pass-by", 6, 120, "Yes"),
    ("chemist_paani", "Paani peeo ji, garmi bahut ae. Te mirchan ghatt khaao!", "F", "chemist", "pass-by", 6, 120, "Yes"),
    ("ustad_swagat", "Akhade vich swagat ae, pehlwan! Dekhde haan kinna dum ae!", "G", "ustad", "dangal-start", 15, 0, "Yes"),
    ("greet_m1", "Sat Sri Akal ji! Ki haal ne?", "F", "any-man", "meet-reply", 6, 20, "Yes"),
    ("greet_young", "Oye, ki haal aa bai? Chadhdi kala!", "E", "young-man", "meet-reply", 6, 20, "Yes"),
    ("greet_old", "Jeonde raho puttar! Sab sukh-saand?", "G", "any-elder", "meet-reply", 6, 20, "Yes"),
    ("greet_f1", "Sat Sri Akal! Aaj badi jaldi vich ae?", "C", "any-woman", "meet-reply", 6, 20, "Yes"),
    ("greet_f2", "Haanji, ki gall ae? Theek-thaak?", "B", "any-woman", "meet-reply", 6, 20, "Yes"),
    ("shop_aao", "Aao bhenji, aao! Ajj sab kujh sasta ae!", "C", "shopkeeper-woman", "pass-by", 6, 40, "Yes"),
    ("shop_aao_m", "Aao ji, aao! Ki chahida ae? Sab kujh milda ethe!", "S", "shopkeeper-man", "pass-by", 6, 40, "Yes"),
    ("player_ssa_m", "Sat Sri Akal ji! Tusi theek ho?", "P", "player", "meet", 6, 8, "Yes"),
    ("player_sapp", "Oye! Aaya maza, sapp! Hahaha!", "R", "player", "snake-hit", 30, 90, "Yes"),
    ("player_ssa_f", "Sat Sri Akal ji! Tusi theek ho?", "Q", "player-f", "meet", 6, 8, "Yes"),
    ("player_uncle_gitte", "Oh uncle ji, tuhada dimaag thode gitteyan vich hai ki?", "P", "player", "fight-elder", 20, 90, "Yes"),
    ("rain_suit", "Haaye mera suit! Chhatri vi ghar reh gayi, hun taan poori bhijj gayi!", "C", "any-woman", "rain-start", 18, 120, "Yes"),
    ("rain_kapde", "Ni Simran, chhatt ton kapde laah chheti! Meeh aa gaya!", "B", "any-woman", "rain-start", 18, 120, "Yes"),
    ("rain_chhatri_aunty", "Hatt ke chal puttar, chhatri naal akh na kadh dayin kise di!", "B", "any-woman, any-elder", "rain-start, flood", 12, 120, "Yes"),
    ("rain_pakode", "Meeh vich taan garam pakode te chah hi bandi ae!", "A", "any-man", "rain-start", 18, 120, "Yes"),
    ("flood_mla", "Har saal eho haal! MLA sahib vote lain aunde ne, naali saaf karaun koi nahi aunda!", "G", "any-elder, any-man", "flood", 18, 120, "Yes"),
    ("flood_naali", "Naaliyan band ne, saara paani gali vich! Committee waale so rahe ne!", "F", "any-man", "flood", 18, 120, "Yes"),
    ("flood_auto_engine", "Paani bahut ae ji, engine band ho jaau! Kiraya doona lagga!", "E", "auto-wala", "flood", 20, 120, "Yes"),
    ("flood_auto_kishti", "Na ji, Dal Bazaar nahi jaana! Othe taan kishti chahidi ae!", "F", "auto-wala", "flood", 20, 120, "Yes"),
    ("gossip_a", "Pata ae? Gillan di nooh Canada chali gayi, te munda ethe hi reh gaya!", "C", "any-woman", "chat", 10, 60, "Yes"),
    ("gossip_b", "Haaye! Main taan pehlan hi kehndi si!", "B", "any-woman", "chat", 10, 60, "Yes"),
    ("chat_match", "Kal da match vekhya? Saadi team ne phir dil tod ditta!", "F", "any-man", "chat", 10, 60, "Yes"),
    ("chat_petrol", "Petrol de bhaa vekh! Hun taan saikal hi changi ae.", "G", "any-man, any-elder", "chat", 10, 60, "Yes"),
    ("call_sabzi", "Aaloo lai lo, pyaaz lai lo! Taaze tamatar, sasti bhindi!", "A", "sabzi-wala", "street-call", 18, 40, "Yes"),
    ("call_fruit", "Kele lai lo, kinnow, mitthe ambh! Chakh ke dekho ji!", "G", "fruit-wala", "street-call", 18, 40, "Yes"),
    ("call_chaat", "Gol gappe! Teekhe vi, mitthe vi! Aao ji, aao!", "A", "chaat-wala", "street-call", 18, 40, "Yes"),
    ("call_doodh", "Doodh waala aa gaya! Taaza doodh, bhenji!", "A", "doodh-wala", "street-call", 20, 40, "Yes"),
    ("call_kabadi", "Kabaadi waala! Purana loha, akhbaar, botlan — vech lo!", "D", "kabadi-wala", "street-call", 20, 40, "Yes"),
    ("call_auto_gill", "Gill, Gill, Gill! Gill Road, Gill! Aa jao, baith jao!", "E", "auto-wala", "street-call", 20, 40, "Yes"),
    ("call_auto_ludhiana", "Ludhiana, Ludhiana, Ludhiana! Bas do sawari hor, chalo chalo!", "F", "auto-wala", "street-call", 20, 40, "Yes"),
    ("call_bus", "Jalandhar, Phagwara, Ludhiana! Chhetti chadho, bus chalan lagi ae!", "D", "bus-conductor, conductor", "street-call", 22, 40, "Yes"),
    ("police_rukk", "Oye! Rukk ja! Rukk ja, hero!", "D", "police, police-bike", "police-chase, police-patrol", 30, 25, "Yes"),
    ("police_chai", "Dekh puttar, case bada serious ae… chai-paani da intezaam karo, te gal khatam!", "D", "police, balwant", "police-bribe", 15, 0, "Yes"),
    ("police_gaddi_rok", "Gaddi rok! Side te laa! Oye, sunda nahi?!", "D", "police-jeep, police-bike", "police-jeep, police-patrol", 45, 12, "Yes"),
    ("lafanga_chai", "Ki dekh reha ae, hero? Saanu chai pila!", "E", "lafanga", "lafanga-taunt", 10, 30, "Yes"),
    ("panga_bhen", "Teri bhen di! Mere naal panga?", "?", "lafanga, any-man", "gang-fight, hit, fight-start", 10, 8, "Yes"),
    ("khad_dasda", "Khad, tainu dasda haan!", "?", "any-man, driver", "blame-both, road-rage, hit, fight-start", 10, 8, "Yes"),
    ("ki_ho_gaya", "Ki ho gaya ji? Ki ho gaya?", "?", "onlooker", "quarrel-onlooker, fight-onlooker", 12, 10, "Yes"),
    ("quarrel_kaun", "Tu jaanda nahi main kaun haan?!", "F", "any-man", "quarrel-start, colony-fight-start", 15, 30, "Yes"),
    ("rage_andha", "Oye! Andha ae? Meri gaddi da nuksaan kaun bharega?!", "F", "driver", "road-rage, crash-drivers", 15, 20, "Yes"),
    ("aunty_fight_1", "Ni kalmuhiye! Tera kutta roz mere boohe agge gand paunda ae!", "C", "aunty", "aunty-quarrel-start, colony-fight-start", 15, 30, "Yes"),
    ("aunty_fight_2", "Chal ni chudail! Pehlan apne vehre di safai kar, phir gall kari!", "B", "aunty", "aunty-quarrel-blows, colony-fight-blows", 15, 30, "Yes"),
    ("fuddu", "Oye, fuddu ho gaya ae ki? Akkal ghaah charan gayi ae!", "E", "passenger", "auto-tipped", 30, 6, "Yes"),
    ("bump_oye", "Oye hoye! Dekh ke chala, hero! This is a bazaar, not a racetrack!", "B", "passenger", "auto-tipped", 30, 6, "Yes"),
    ("safari_kanjar", "Oye kanjar da puttar! Auto palat ditta! Hun paise nahi dene tere!", "F", "passenger", "auto-tipped", 30, 6, "Yes"),
    ("sawari_lak", "Hai rabba! Mera lakk tutt gaya! Kithe dhyaan ae tera, fuddu?!", "C", "passenger", "auto-tipped", 30, 6, "Yes"),
    # Trial batch: 50 new lines to hear how the voices and lines work in the game (see docs/VOICE_LINES.md).
    ("trial_meet_01", "Sat Sri Akal ji! Sab theek-thaak?", "A", "any-man", "meet-reply", 7, 8, "Yes"),
    ("trial_meet_02", "Ki haal ae puttar? Chadhdi kala!", "G", "any-elder", "meet-reply", 7, 8, "Yes"),
    ("trial_meet_03", "Haanji bhenji, kiddan ho? Sab khair ae?", "C", "any-woman", "meet-reply", 7, 8, "Yes"),
    ("trial_meet_04", "Oye veere, bahut dinaan baad disda!", "E", "young-man", "meet-reply", 7, 8, "Yes"),
    ("trial_meet_05", "Jeonda reh puttar! Aaj bazaar kidhar nu?", "G", "any-elder", "meet-reply", 7, 8, "Yes"),
    ("trial_meet_06", "Ni, kinne din ho gaye, dikhya nahi!", "T", "any-woman", "meet-reply", 7, 8, "Yes"),
    ("trial_meet_07", "Aao ji, Sat Sri Akal! Kuchh lainde ho?", "S", "shopkeeper-man", "meet-reply", 7, 8, "Yes"),
    ("trial_meet_08", "Kiddan bhenji? Ghar de sab theek?", "T", "any-woman", "meet-reply", 7, 8, "Yes"),
    ("trial_meet_09", "Sat Sri Akal ji, Waheguru di kirpa ae!", "A", "any-man", "meet-reply", 7, 8, "Yes"),
    ("trial_meet_10", "Bas ji, chal reha ae. Tu sunao!", "E", "young-man", "meet-reply", 7, 8, "Yes"),
    ("trial_meet_11", "Haanji bhenji, sab theek, tusi sunao!", "C", "any-woman", "meet-reply", 7, 8, "Yes"),
    ("trial_meet_12", "Waheguru mehar kare, bas chal reha ae!", "G", "any-elder", "meet-reply", 7, 8, "Yes"),
    ("trial_pass_01", "Dhyaan naal, gali vich paani khaloya ae!", "B", "any-woman", "flood", 18, 120, "Yes"),
    ("trial_pass_02", "Sabzi da bhaa vekh, ajj taan hosh ud gaye!", "C", "shopkeeper-woman", "pass-by", 9, 90, "Yes"),
    ("trial_pass_03", "Hai Waheguru, eh garmi kadon jaave!", "T", "shopkeeper-woman", "pass-by", 9, 90, "Yes"),
    ("trial_pass_04", "Petrol vi mehnga, zindagi vi mehngi ho gayi!", "A", "shopkeeper-man", "pass-by", 9, 90, "Yes"),
    ("trial_pass_05", "Bijli phir chali gayi! Inverter vi thak gaya!", "F", "shopkeeper-man", "pass-by", 9, 90, "Yes"),
    ("trial_pass_06", "Oye, ghar aaiye, chai taiyaar ae!", "G", "shopkeeper-man", "pass-by", 9, 90, "Yes"),
    ("trial_pass_07", "Eh sadak taan hamesha tootti rehndi ae, kaun sunda ae!", "F", "shopkeeper-man", "pass-by", 9, 90, "Yes"),
    ("trial_pass_08", "Ajj taan kanak da bhaa vadh gaya ae, sun ke dil ghabraaunda!", "G", "shopkeeper-man", "pass-by", 9, 90, "Yes"),
    ("trial_pass_09", "Bhai, bike di chain dheeli ae, dhyaan karo!", "E", "shopkeeper-man", "pass-by", 9, 90, "Yes"),
    ("trial_pass_10", "Veere, gaddi sambhal ke chala, sadak te bachche vi ne!", "F", "shopkeeper-man", "pass-by", 9, 90, "Yes"),
    ("trial_pass_11", "Waheguru ji, bachche kithe ne? Dhyaan rakho!", "C", "shopkeeper-woman", "pass-by", 9, 90, "Yes"),
    ("trial_pass_12", "Mausam vi ajeeb ae, subah thand, dopahir garmi!", "B", "any-woman", "rain-start", 18, 120, "Yes"),
    ("trial_pass_13", "Chai pee le, garam garam! Bahut sohni ae!", "A", "shopkeeper-man", "pass-by", 9, 90, "Yes"),
    ("trial_chat_01", "Oye, Bhola da dhaba mehenga ho gaya ae ajkal!", "A", "any-man", "chat", 12, 90, "No"),
    ("trial_chat_02", "Waheguru, mausam kadon badlega? Ik vaari thand aave!", "B", "any-woman", "chat", 12, 90, "No"),
    ("trial_chat_03", "Mere munde ne naya phone lia, hun taan udda phirda ae!", "C", "any-woman", "chat", 12, 90, "No"),
    ("trial_call_01", "Chaat, chaat! Teekhi meethi, khaa ke dekho ji!", "A", "chaat-wala", "street-call", 18, 90, "Yes"),
    ("trial_call_02", "Ganne da ras! Taaza, thanda thanda, dus rupay gilaas!", "C", "fruit-wala", "street-call", 18, 90, "Yes"),
    ("trial_call_03", "Gobhi, matar, aloo! Taaze taaze, sasti sabzi!", "S", "sabzi-wala", "street-call", 18, 90, "Yes"),
    ("trial_call_04", "Auto! Auto! Ghanta Ghar, Chaura, Gill Road, aa jao!", "E", "auto-wala", "street-call", 20, 90, "Yes"),
    ("trial_call_05", "Nawan Shehar, Nawan Shehar! Ik sawari hor, chalo!", "F", "auto-wala", "street-call", 20, 90, "Yes"),
    ("trial_call_06", "Naya suit, sasta suit! Aao bhenji, aao!", "U", "shopkeeper-woman", "street-call", 15, 90, "Yes"),
    ("trial_call_07", "Samosa, samosa! Garam samosa, do ke teen!", "S", "shopkeeper-man", "street-call", 15, 90, "Yes"),
    ("trial_call_08", "Kulfi! Matka kulfi, kamaal di! Aao ji!", "A", "shopkeeper-man", "street-call", 15, 90, "Yes"),
    ("trial_call_09", "Dahi, makhan, lassi! Thandi lassi, pind wali!", "E", "shopkeeper-man", "street-call", 15, 90, "Yes"),
    ("trial_call_10", "Aaj da akhbaar, paanch rupay! Taaza, taaza!", "E", "anyone", "street-call", 15, 90, "Yes"),
    ("trial_call_11", "Chhole kulche, garma-garam! Ik plate da das rupay!", "A", "shopkeeper-man", "street-call", 15, 90, "Yes"),
    ("trial_call_12", "Pakode, pakode! Garam pakode, khatte chutney naal!", "S", "shopkeeper-man", "street-call", 15, 90, "Yes"),
    ("trial_police_01", "Gaddi side te laa! Licence te kaagaz dikha!", "D", "police-jeep, police-bike", "police-patrol", 25, 90, "Yes"),
    ("trial_police_02", "Oye, gaddi ikk pase kar, rasta chhad!", "D", "police", "police-chase", 20, 90, "Yes"),
    ("trial_police_03", "Challo bhai, kaagaz poore nahi, chai-paani karo!", "D", "police", "police-bribe", 8, 90, "Yes"),
    ("trial_police_04", "Gaddi rok ke, naam te pata daso! Hun bhajna nahi!", "D", "police", "police-chase", 20, 90, "Yes"),
    ("trial_police_05", "Sarkar di gaddi! Side te ho jao, sunda nahi?!", "D", "police-jeep", "police-jeep", 30, 90, "Yes"),
    ("trial_police_06", "Heroagiri nahi chalegi yahan! Ruk ja!", "D", "police", "police-chase", 20, 90, "Yes"),
    ("trial_fight_01", "Tu mere naal gall na kar, main jaanda haan tenu!", "F", "any-man", "quarrel-start", 15, 90, "Yes"),
    ("trial_fight_02", "Chal aaja, dekh lainde haan ik vaari!", "E", "young-man", "fight-start", 12, 90, "Yes"),
    ("trial_fight_03", "Bas kar ji, saare mohalle da tamasha na banaao!", "T", "any-woman", "quarrel-onlooker", 14, 90, "Yes"),
    ("trial_fight_04", "Oye, rehan de, ghar jaake gall karo!", "G", "any-elder", "quarrel-onlooker", 14, 90, "Yes"),
    # Session 14: 50 lines for the doorbell, power cuts, the level crossing, the pump, bank, post, signals and friends.
    ("v2_bell_01", "Kaun ae, ji? Ki gal ae?", "T", "any-woman, any-elder", "doorbell-stranger", 12, 20, "Yes"),
    ("v2_bell_02", "Haanji, ki chahida? Ruko, aa rahi haan!", "C", "any-woman", "doorbell-stranger", 12, 20, "Yes"),
    ("v2_bell_03", "Sadde ghar kuj nahi chahida, ji. Shukriya!", "T", "any-woman, any-elder", "doorbell-stranger", 12, 20, "Yes"),
    ("v2_bell_04", "Haanji paaji, ki hukam ae?", "A", "any-man", "doorbell-stranger", 12, 20, "Yes"),
    ("v2_bell_05", "Kiss nu milna ae, puttar? Andar aa ja!", "S", "any-man, any-elder", "doorbell-stranger", 12, 20, "Yes"),
    ("v2_bell_06", "Oye kaun ae? Zara ruk, buddhe nu uthan de!", "G", "any-man, any-elder", "doorbell-stranger", 12, 20, "Yes"),
    ("v2_bell_07", "Aa gaya puttar? Waheguru, bahut din baad!", "G", "any-man, any-elder", "doorbell-family", 12, 20, "Yes"),
    ("v2_bell_08", "Hai, galat ghar aa gaye ho shayad, paaji.", "B", "any-woman, any-elder", "doorbell-stranger", 12, 20, "Yes"),
    ("v2_gate_01", "Oye, tu aa gaya! Andar aa, chah pee ke jaayi!", "A", "any-man", "doorbell-family", 12, 20, "Yes"),
    ("v2_gate_02", "Haye puttar, bahut vadhia kita jo aa gaya! Mummy ghar ae.", "G", "any-man, any-elder", "doorbell-family", 12, 20, "Yes"),
    ("v2_gate_03", "Sat Sri Akal ji! Aaj bahut dina baad chakkar laaya ae!", "T", "any-woman, any-elder", "doorbell-family", 12, 20, "Yes"),
    ("v2_gate_04", "Aaja aaja, gate khula ae! Pehchaan liya si tenu!", "C", "any-woman", "doorbell-family", 12, 20, "Yes"),
    ("v2_night_01", "Eni raat nu kaun ae? Subah aana, puttar!", "G", "any-elder", "doorbell-night", 20, 20, "Yes"),
    ("v2_night_02", "Raat bahut ho gayi, kal aao ji!", "T", "any-woman", "doorbell-night", 20, 20, "Yes"),
    ("v2_night_03", "Oye! Sab sau gaye ne, subah gal karange!", "F", "any-man", "doorbell-night", 20, 20, "Yes"),
    ("v2_siesta_01", "Oh ho, dupehar nu sau rahe haan! Baad vich aana.", "B", "any-woman", "doorbell-siesta", 20, 20, "Yes"),
    ("v2_siesta_02", "Hai rabba, ajj vi araam nahi! Kaun ae?", "T", "any-woman", "doorbell-siesta", 20, 20, "Yes"),
    ("v2_cut_01", "Hai rabba, bijli phir gayi! Inverter laga lo!", "C", "any-woman", "power-cut", 20, 120, "Yes"),
    ("v2_cut_02", "PSPCL walon ne phir kamaal kar ditta!", "A", "any-man", "power-cut", 20, 120, "Yes"),
    ("v2_cut_03", "Roz roz light cut, hun taan hadd ho gayi!", "F", "any-man", "power-cut", 20, 120, "Yes"),
    ("v2_cut_04", "Mobile da charge vi nahi hoya si, haaye!", "T", "any-woman", "power-cut", 20, 120, "Yes"),
    ("v2_cut_05", "Saade vele taan lalten hundi si, koi fikar nahi!", "G", "any-elder", "power-cut", 20, 120, "Yes"),
    ("v2_back_01", "Aa gayi bijli! Shukar ae Waheguru da!", "E", "any-man, young-man", "power-back", 20, 120, "Yes"),
    ("v2_back_02", "Light aa gayi! Chalo, tv lagao!", "C", "any-woman", "power-back", 20, 120, "Yes"),
    ("v2_rail_01", "Phatak band ho gaya! Hun das minute!", "F", "any-man", "rail-gate", 25, 120, "Yes"),
    ("v2_rail_02", "Train aa rahi ae, ruko ji, ruko!", "A", "any-man", "rail-gate", 25, 120, "Yes"),
    ("v2_rail_03", "Eh phatak taan roz hi band rehnda ae!", "T", "any-woman", "rail-gate", 25, 120, "Yes"),
    ("v2_rail_04", "Saade vele railway line te koi gaddi nahi aundi si.", "G", "any-elder", "rail-gate", 25, 120, "Yes"),
    ("v2_pump_01", "Full kar ditta, paaji! Safar mubarak!", "A", "pumpwala", "pump-done", 15, 30, "Yes"),
    ("v2_pump_02", "Petrol, diesel, hawa; sab kujh milda ethe!", "S", "pumpwala", "pump-arrive", 20, 120, "Yes"),
    ("v2_pump_03", "Meter dekh lo ji, zero ton shuru kitta si!", "F", "pumpwala", "pump-done", 15, 30, "Yes"),
    ("v2_fuel_p1", "Oye! Tank khaali ho gaya, hun dhakka lagauna painda!", "P", "player", "fuel-out", 30, 120, "Yes"),
    ("v2_fuel_p2", "Hai rabba, petrol khatam! Pump kithe ae?", "P", "player", "fuel-out", 30, 120, "Yes"),
    ("v2_fuel_q1", "Hai rabba, tank khaali! Hun ki karaan?", "Q", "player-f", "fuel-out", 30, 120, "Yes"),
    ("v2_fuel_q2", "Oho, petrol khatam! Koi dhakka laga do!", "Q", "player-f", "fuel-out", 30, 120, "Yes"),
    ("v2_bank_01", "Aao ji, Panj Aab Bank vich swagat ae. Ki seva karaan?", "S", "clerk_bank", "bank-welcome", 15, 120, "Yes"),
    ("v2_bank_02", "Paise bank vich rakho, jeb vich nahi, puttar.", "S", "clerk_bank", "bank-advice", 15, 120, "Yes"),
    ("v2_post_01", "Aao ji, ajj di dak aayi ae tuhadi.", "G", "clerk_post", "post-welcome", 15, 120, "Yes"),
    ("v2_post_02", "Chitthi aayi ae, Canada ton! Eh lo ji.", "G", "clerk_post", "post-welcome", 15, 120, "Yes"),
    ("v2_signal_01", "Oye lal batti si! Akhan band ne ki?!", "F", "any-man", "signal-jump", 25, 90, "Yes"),
    ("v2_signal_02", "Veere, signal tod ke kitthe chaleya?", "E", "any-man, young-man", "signal-jump", 25, 90, "Yes"),
    ("v2_signal_03", "Hai, hun taan thok dinda ae! Dhyaan naal!", "B", "any-woman", "signal-jump", 25, 90, "Yes"),
    ("v2_friend_01", "Oye sher! Kidhar gayab ho gaya si, yaar?", "E", "any-man, young-man", "meet-friend", 7, 20, "Yes"),
    ("v2_friend_02", "Chal, chai pee lainde haan, sab kujh daseyo!", "F", "any-man", "meet-friend", 7, 20, "Yes"),
    ("v2_friend_03", "Yaar, tu taan hun bada banda ban gaya ae!", "A", "any-man", "meet-friend", 7, 20, "Yes"),
    ("v2_friend_04", "Tere bina taan gali sooni lagdi si, veere!", "B", "any-woman", "meet-friend", 7, 20, "Yes"),
    ("v2_meet_01", "Ae ki, tu Gurnam da munda nahi? Kinna vadda ho gaya!", "S", "any-man", "meet-reply", 7, 20, "No"),
    ("v2_meet_02", "Nava gaddi? Waah, chhakka maar ditta tu taan!", "D", "any-man, young-man", "vehicle-admire", 12, 60, "Yes"),
    ("v2_meet_03", "Bhenji, tuhada puttar taan bade kamaal da ae!", "Q", "any-woman", "chat", 12, 90, "Yes"),
    ("v2_meet_04", "Sat Sri Akal, ji! Kaim ho? Pind vich sab changa?", "P", "player", "meet", 6, 8, "Yes"),
    ("v3_chaupal_01", "Pehlan shaam nu chaupal te raunak hundi si, hun koi nahi aunda.", "G", "any-elder", "chat", 14, 90, "Yes"),
    ("v3_chaupal_02", "Sun puttar, gallan karan naal pet nahi bharda!", "G", "any-elder", "chat", 14, 90, "Yes"),
    ("v3_chaupal_03", "Charpai te dhoop sekhde haan, hun chah pila de.", "G", "any-elder", "chat", 14, 90, "Yes"),
    ("v3_chaupal_04", "Saadi umar vich taan bailgaddi si, eh gaddiyan kithon aa gayiyan!", "G", "any-elder", "chat", 14, 90, "Yes"),
    ("v3_chaupal_05", "Mainu kujh na pucho, main taan sab pehlan hi jaanda haan!", "G", "any-elder", "chat", 14, 90, "Yes"),
    ("v3_gossip_01", "Ni, Sharma ji da munda vi Canada jaa reha ae, sach!", "T", "aunty", "chat", 14, 90, "Yes"),
    ("v3_gossip_02", "Haaye, ohde ghar Mercedes aayi ae, dekhi ni?", "T", "aunty", "chat", 14, 90, "Yes"),
    ("v3_gossip_03", "Bijli gayi, tv vi band, hun gallan hi karo ni!", "T", "aunty", "chat", 14, 90, "Yes"),
    ("v3_gossip_05", "Mere vale ne aaj sabzi mehngi dasi, sach ni!", "T", "aunty", "chat", 14, 90, "Yes"),
    ("v3_pass_e_01", "Oye puttar, gaddi dheere chala, bachche ne gali vich!", "G", "any-elder", "chat", 12, 90, "Yes"),
    ("v3_pass_e_02", "Eh dhuan kadon ghatt hoega, kaun sunda ae!", "G", "any-elder", "chat", 12, 90, "Yes"),
    ("v3_pass_e_03", "Chal chal, raah chhad, budhe nu langhan de!", "G", "any-elder", "chat", 12, 90, "Yes"),
    ("v3_pass_w_01", "Hai rabba, eh dhool, kapde saare mitti ho gaye!", "C", "any-woman", "chat", 12, 90, "Yes"),
    ("v3_pass_w_03", "Oye, eh rickshaw wala vi raah nahi dekhda!", "C", "any-woman", "chat", 12, 90, "Yes"),
    # Elders answer a greeting ("meet-reply"); "talk" is reserved for named dialogue, so a passer-by never reaches these.
    ("v3_talk_e_01", "Puttar, tera naam kiddan? Ikk vaar dass, yaad rakhangey!", "G", "any-elder", "meet-reply", 7, 20, "Yes"),
    ("v3_talk_e_02", "Dhan waheguru, tu ajj ghar aaya te chah pee ke jaa!", "G", "any-elder", "meet-reply", 7, 20, "Yes"),
    # v3 batch 2: bumps, auto jolts, quarrels, colony blows, a Safari, and four more greeting replies.
    ("v3_bump_01", "Oye! Andha ae ke? Dikhda nahi tainu?", "F", "any-man", "bumped", 8, 90, "Yes"),
    ("v3_bump_02", "Hai rabba, mari ja rahi si! Dhyaan naal chala!", "T", "any-woman", "bumped", 8, 90, "Yes"),
    ("v3_bump_03", "Puttar, budhe nu takkar maar ditti! Sharam kar!", "G", "any-elder", "bumped", 8, 90, "Yes"),
    ("v3_bump_04", "Bhai, bike mod ke chala, sadak sadi vi ae!", "E", "young-man", "bumped", 8, 90, "Yes"),
    ("v3_jolt_01", "Hai rabba, ki khadda ae! Haddi-pasli hil gayi!", "C", "passenger", "ride-jolt", 6, 90, "Yes"),
    ("v3_jolt_02", "Auto wale, ik vaar dheere chala, main utar jaana!", "A", "passenger", "ride-jolt", 6, 90, "Yes"),
    ("v3_jolt_03", "Oye, khadde vich gaddi ghusa ke kithe chali?!", "E", "passenger", "ride-jolt", 6, 90, "Yes"),
    ("v3_qblow_01", "Oye hatt ja! Hun haath chal gaye ne!", "F", "any-man", "quarrel-blows", 14, 90, "Yes"),
    ("v3_qblow_02", "Lo ji, hun taan maar-kutt shuru ho gayi!", "E", "any-man", "quarrel-blows", 14, 90, "Yes"),
    ("v3_qblow_03", "Chhadd ke, bhai! Mohalle wale aa gaye ne!", "A", "any-man", "quarrel-blows", 14, 90, "Yes"),
    ("v3_cfight_01", "Bas kar ji, mohalle da naam kharab na kar!", "T", "any-woman", "colony-fight-blows", 14, 90, "Yes"),
    ("v3_cfight_02", "Puttar, ghar jaa ke gall kar, gali vich maar-kutt na!", "G", "any-elder", "colony-fight-blows", 14, 90, "Yes"),
    ("v3_safari_01", "Oye! Meri safari nu thok ditta! Hun paisa de!", "F", "driver", "safari-hit", 10, 90, "Yes"),
    ("v3_safari_02", "Kanjar da put, gaddi di kadar nahi? Maafi mang!", "E", "driver", "safari-hit", 10, 90, "Yes"),
    ("v3_talk_g_01", "Ki haal ae puttar? Dass, sab khair ae?", "G", "any-elder", "meet-reply", 7, 20, "Yes"),
    ("v3_talk_g_02", "Bhenji, ajj bazaar vich bhaddi ae, kiddan ho?", "C", "any-woman", "meet-reply", 7, 20, "Yes"),
    ("v3_talk_g_03", "Oye veere, kiddan? Chai peeni ae taan dasso!", "A", "any-man", "meet-reply", 7, 20, "Yes"),
    ("v3_talk_g_04", "Ki hal chal ae, pind vich sab changa?", "G", "any-elder", "meet-reply", 7, 20, "Yes"),
    # v3 batch 3: two drivers after a crash, a passer-by admiring a new car, and a lafanga taunting the player.
    ("v3_crash_01", "Oye, gaddi nu thokke maar ke kithe chal reha ae? Hun mere naal gall kar!", "F", "driver", "crash-drivers", 15, 90, "Yes"),
    ("v3_crash_02", "Mera bumper tutt gaya! Hun paisa de, nahi taan thana chalega!", "D", "driver", "crash-drivers", 15, 90, "Yes"),
    ("v3_admire_01", "Wah ji, kinna sohna gaddi ae! Kine layi ae?", "E", "any-man, young-man", "vehicle-admire", 12, 60, "Yes"),
    ("v3_admire_02", "Oye, nava model! Pind vich eda kise kol nahi!", "A", "any-man, young-man", "vehicle-admire", 12, 60, "Yes"),
    ("v3_lafanga_01", "Oye hero, ajj kal bahut akad ho gayi ae? Sadi gali vich aa ke dekh!", "E", "lafanga", "lafanga-taunt", 10, 90, "Yes"),
    ("v3_lafanga_02", "Teri gaddi taan bahut chamkdi ae, hero! Sadi taraf dhyaan rakhi!", "E", "lafanga", "lafanga-taunt", 10, 90, "Yes"),
]

HEAD = ["#", "Clip file", "What it says", "Voice", "Play it? (Yes/No)", "Who says it", "When it plays", "Max distance (m)", "Min gap before repeat (s)", "Your notes"]

def codes(s):
    return [c.strip() for c in str(s or "").replace(";", ",").split(",") if c.strip()]

def is_greet(clip, when):
    """Greetings may repeat on their own short gap; every other clip waits at least 90 s."""
    return bool({"meet", "meet-reply", "meet-friend", "vehicle-admire"} & set(when)) or any(w.startswith("doorbell") for w in when) or clip.startswith("shop_aao")

def write_json(rows, path="src/audio/voiceMap.json"):
    out = {r[0]: {"on": str(r[7]).strip().lower().startswith("y"), "who": codes(r[3]), "when": codes(r[4]), "maxDist": float(r[5] or 15), "gap": float(r[6] or 0) if is_greet(r[0], codes(r[4])) else max(90.0, float(r[6] or 0)), "greet": is_greet(r[0], codes(r[4])), "say": str(r[1] or ""), "voice": str(r[2] or "?")[:1]} for r in rows}
    with open(path, "w") as f: json.dump(out, f, indent=1, ensure_ascii=False)
    print("wrote", path, len(out), "clips")

def make():
    wb = Workbook()
    ws = wb.active; ws.title = "Voice clips"
    ws.append(HEAD)
    for i, (clip, words, voice, who, when, dist, gap, on) in enumerate(CLIPS, 1):
        ws.append([i, clip, words, f"{voice} — {VOICES[voice]}", on, who, when, dist, gap if is_greet(clip, codes(when)) else max(90, gap), ""])
    widths = [4, 20, 62, 30, 12, 30, 34, 12, 14, 30]
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    hdr = PatternFill("solid", fgColor="E8A33C"); edit = PatternFill("solid", fgColor="FFF6DA")
    for c in ws[1]: c.font = Font(bold=True); c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[4:10]: c.fill = edit
    ws.freeze_panes = "C2"
    n = len(CLIPS) + 1
    yn = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False); ws.add_data_validation(yn); yn.add(f"E2:E{n}")
    # Who/When accept several codes separated by commas, so the lists are a help, not a lock.
    wv = DataValidation(type="list", formula1=f"Codes!$A$2:$A${len(WHO) + 1}", showErrorMessage=False); ws.add_data_validation(wv); wv.add(f"F2:F{n}")
    tv = DataValidation(type="list", formula1=f"Codes!$C$2:$C${len(WHEN) + 1}", showErrorMessage=False); ws.add_data_validation(tv); tv.add(f"G2:G{n}")
    cs = wb.create_sheet("Codes")
    cs.append(["Who code", "Who it means", "When code", "When it means"])
    for i in range(max(len(WHO), len(WHEN))):
        a = WHO[i] if i < len(WHO) else ("", ""); b = WHEN[i] if i < len(WHEN) else ("", "")
        cs.append([a[0], a[1], b[0], b[1]])
    for c in cs[1]: c.font = Font(bold=True); c.fill = hdr
    for col, w in zip("ABCD", [18, 46, 22, 48]): cs.column_dimensions[col].width = w
    hs = wb.create_sheet("How to fill")
    for line in [
        "Edit only the yellow columns on the 'Voice clips' sheet.",
        "Play it?: No = this clip never plays.",
        "Who says it: one or more codes from the Codes sheet, separated by commas (e.g. 'gill, any-elder').",
        "When it plays: one or more codes from the Codes sheet, separated by commas (e.g. 'pass-by, talk').",
        "Max distance: the clip only plays if that person is within this many metres of you (5-8 for greetings).",
        "Min gap: seconds before the same clip can play again (never less than 90, except greetings).",
        "A clip whose Who does not match the person speaking, or whose When does not match the moment, stays silent.",
        "Named characters (Bhola, Gurnam, Pammi, Masterji, ...) say ONLY clips that list them by name - group codes like any-woman never apply to them.",
        "Afterwards send the file back (or run: python3 tools/voice_map.py import docs/voice-map.xlsx) and the game follows it.",
        "Several clips may share a Who + When; then the game picks one of them without repeating the last.",
        "Voices: several characters share one recorded voice (see the Voice column). Ask for new takes to give someone their own.",
    ]: hs.append([line])
    hs.column_dimensions["A"].width = 120
    wb.save("docs/voice-map.xlsx"); print("wrote docs/voice-map.xlsx")
    write_json([(c[0], c[1], c[2], c[3], c[4], c[5], c[6], c[7]) for c in CLIPS])

def imp(path):
    ws = load_workbook(path, data_only=True)["Voice clips"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[1]: continue
        rows.append((r[1], r[2], r[3], r[5], r[6], r[7], r[8], r[4]))
    write_json(rows)

if __name__ == "__main__":
    make() if sys.argv[1:2] == ["make"] else imp(sys.argv[2])
